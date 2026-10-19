"""Excel workbook for each run: one filterable row per posting, plus a summary sheet.

Sheets
  Matches  - one row per reported posting, best first, with a clickable title,
             monthly salary as numbers, gaps, and blank Status / Notes columns
             for tracking applications.
  Summary  - run date, tier thresholds and counts. Counts are COUNTIF formulas
             over the Matches sheet, so they update if you delete or re-tier rows.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import GradedJob
from .report import TIER_LABEL, _level, source_label

FONT = "Arial"
TIER_FILL = {"A": "D4EDDA", "B": "DCEBFA", "C": "FFF3CD", "D": "EEEEEE"}
STATUS_CHOICES = "To review,Shortlisted,Applied,Interviewing,Offer,Rejected,Not interested"

COLUMNS = [
    # header, width
    ("Tier", 6),
    ("Requirements met", 12),
    ("Score %", 9),
    ("AI score", 9),
    ("Title", 46),
    ("Company", 30),
    ("Salary min (SGD/mo)", 12),
    ("Salary max (SGD/mo)", 12),
    ("Salary as posted", 22),
    ("Level", 16),
    ("Arrangement", 12),
    ("Employment type", 18),
    ("Source", 15),
    ("Posted", 11),
    ("Gaps (requirements not met)", 48),
    ("AI note", 40),
    ("Link", 40),
    ("Status", 14),
    ("Notes", 30),
]
COL = {name: i + 1 for i, (name, _) in enumerate(COLUMNS)}


def build_workbook(jobs: list[GradedJob], run_date: datetime, stats: dict, tiers: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Matches"
    _write_matches(ws, jobs)
    _write_summary(wb.create_sheet("Summary"), len(jobs), run_date, stats, tiers)
    return wb


def workbook_bytes(jobs: list[GradedJob], run_date: datetime, stats: dict, tiers: dict) -> bytes:
    buf = BytesIO()
    build_workbook(jobs, run_date, stats, tiers).save(buf)
    return buf.getvalue()


def write_workbook(path: str | Path, jobs: list[GradedJob], run_date: datetime, stats: dict, tiers: dict) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(workbook_bytes(jobs, run_date, stats, tiers))
    return path


def _write_matches(ws, jobs: list[GradedJob]) -> None:
    header_font = Font(name=FONT, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    body_font = Font(name=FONT, size=10)
    link_font = Font(name=FONT, size=10, color="0563C1", underline="single")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    for i, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    for r, g in enumerate(jobs, start=2):
        j = g.job
        lo, hi = j.monthly_salary()
        values = {
            "Tier": g.tier,
            "Requirements met": f"{g.met_count} of {g.total_count}",
            "Score %": g.score_pct / 100.0,
            "AI score": g.ai_score,
            "Title": j.title,
            "Company": j.company,
            "Salary min (SGD/mo)": round(lo) if lo is not None else None,
            "Salary max (SGD/mo)": round(hi) if hi is not None else None,
            "Salary as posted": j.salary_display(),
            "Level": _level(g),
            "Arrangement": j.work_arrangement or "",
            "Employment type": j.employment_type or "",
            "Source": source_label(j.source),
            "Posted": j.posted_at.replace(tzinfo=None).date() if j.posted_at else None,
            "Gaps (requirements not met)": g.gaps() if g.unmet else "",
            "AI note": g.ai_reason or "",
            "Link": j.url,
            "Status": "To review",
            "Notes": "",
        }
        fill = PatternFill("solid", fgColor=TIER_FILL.get(g.tier, "FFFFFF"))
        for name, value in values.items():
            cell = ws.cell(row=r, column=COL[name], value=value)
            cell.font = body_font
            cell.alignment = wrap if name in ("Title", "Gaps (requirements not met)", "AI note", "Notes") else top
        tier_cell = ws.cell(row=r, column=COL["Tier"])
        tier_cell.fill = fill
        tier_cell.font = Font(name=FONT, size=10, bold=True)
        tier_cell.alignment = Alignment(horizontal="center", vertical="top")
        if j.url:
            for name in ("Title", "Link"):
                c = ws.cell(row=r, column=COL[name])
                c.hyperlink = j.url
                c.font = link_font
        ws.cell(row=r, column=COL["Score %"]).number_format = "0%"
        for name in ("Salary min (SGD/mo)", "Salary max (SGD/mo)"):
            ws.cell(row=r, column=COL[name]).number_format = "#,##0"
        ws.cell(row=r, column=COL["Posted"]).number_format = "dd mmm yyyy"

    last_row = max(2, len(jobs) + 1)
    ws.freeze_panes = "F2"  # keep Tier..Title visible while scrolling right
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"
    dv = DataValidation(type="list", formula1=f'"{STATUS_CHOICES}"', allow_blank=True)
    status_col = get_column_letter(COL["Status"])
    dv.add(f"{status_col}2:{status_col}{max(last_row, 500)}")
    ws.add_data_validation(dv)


def _write_summary(ws, count: int, run_date: datetime, stats: dict, tiers: dict) -> None:
    bold = Font(name=FONT, bold=True)
    body = Font(name=FONT)
    title = Font(name=FONT, bold=True, size=14)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60

    ws["A1"] = f"Job Digest – {run_date.strftime('%a %d %b %Y')} (SGT)"
    ws["A1"].font = title

    tier_col = get_column_letter(COL["Tier"])
    rng = f"Matches!${tier_col}$2:${tier_col}${max(2, count + 1)}"
    status_col = get_column_letter(COL["Status"])
    status_rng = f"Matches!${status_col}$2:${status_col}${max(2, count + 1)}"

    rows = [
        ("Tier", "Postings", "Rule"),
        ("A – strong match", f'=COUNTIF({rng},"A")', f"Score ≥ {tiers.get('A', 80):.0f}%"),
        ("B – good match", f'=COUNTIF({rng},"B")', f"Score ≥ {tiers.get('B', 60):.0f}%"),
        ("C – partial match", f'=COUNTIF({rng},"C")', f"Score ≥ {tiers.get('C', 40):.0f}%"),
        ("Total new matches", "=SUM(B4:B6)", "Sum of the tiers above"),
    ]
    for i, row in enumerate(rows, start=3):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = bold if i in (3, 7) else body
    for i, t in zip(range(4, 7), "ABC"):
        ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=TIER_FILL[t])

    ws["A9"], ws["A9"].font = "Application tracking", bold
    ws["A10"], ws["B10"] = "Applied", f'=COUNTIF({status_rng},"Applied")'
    ws["A11"], ws["B11"] = "Still to review", f'=COUNTIF({status_rng},"To review")'
    for cell in ("A10", "B10", "A11", "B11"):
        ws[cell].font = body

    ws["A13"], ws["A13"].font = "This run (as reported by the search)", bold
    per_source = ", ".join(f"{k} {v}" for k, v in (stats.get("per_source") or {}).items())
    facts = [
        ("Postings fetched", stats.get("fetched", 0), per_source),
        ("Duplicates across boards", stats.get("duplicates", 0), ""),
        ("Excluded by must-have rules", stats.get("excluded", 0), "Wrong title, accounting role, excluded term, location or company"),
        ("Already sent in earlier digests", stats.get("seen", 0), "Not repeated"),
    ]
    for i, (label, value, note) in enumerate(facts, start=14):
        ws.cell(row=i, column=1, value=label).font = body
        ws.cell(row=i, column=2, value=value).font = body
        ws.cell(row=i, column=3, value=note).font = body

    ws["A19"], ws["A19"].font = "How to read the Matches sheet", bold
    notes = [
        "Requirements met counts every requirement equally. Score % is weighted, so the industry requirement (weight 3) counts three times.",
        "Salary columns are converted to SGD per month. Blank means the posting does not state a salary.",
        "Use the Status column drop-down to track applications; the counts above update automatically.",
    ]
    for i, n in enumerate(notes, start=20):
        c = ws.cell(row=i, column=1, value=n)
        c.font = body
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30
