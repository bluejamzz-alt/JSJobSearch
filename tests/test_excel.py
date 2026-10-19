from datetime import datetime
from email import message_from_bytes
from io import BytesIO

from openpyxl import load_workbook

from jobsearch.excel import COL, build_workbook, workbook_bytes
from jobsearch.grader import grade_job
from jobsearch.notify import email as email_mod


def _graded(config, job, **changes):
    for k, v in changes.items():
        setattr(job, k, v)
    return grade_job(job, config["requirements"], config["grading"])


def test_workbook_rows_links_and_salary(config, good_job):
    a = _graded(config, good_job)
    wb = build_workbook([a], datetime(2026, 10, 19, 8), {"fetched": 5, "per_source": {"MyCareersFuture": 5}},
                        config["grading"]["tiers"])
    assert wb.sheetnames == ["Matches", "Summary"]
    ws = wb["Matches"]
    assert ws.cell(1, COL["Title"]).value == "Title"
    assert ws.cell(2, COL["Tier"]).value == "A"
    assert ws.cell(2, COL["Requirements met"]).value == "11 of 11"
    assert ws.cell(2, COL["Score %"]).value == 1.0
    assert ws.cell(2, COL["Title"]).hyperlink.target == "https://example.com/1"
    assert ws.cell(2, COL["Salary min (SGD/mo)"]).value == 4500
    assert ws.cell(2, COL["Salary max (SGD/mo)"]).value == 6000
    assert ws.cell(2, COL["Status"]).value == "To review"
    assert ws.auto_filter.ref.startswith("A1:")
    assert ws.freeze_panes == "F2"


def test_annual_salary_converted_and_missing_salary_blank(config, good_job):
    g = _graded(config, good_job, salary_min=60000, salary_max=72000, salary_period="annual")
    ws = build_workbook([g], datetime(2026, 10, 19), {}, config["grading"]["tiers"])["Matches"]
    assert ws.cell(2, COL["Salary min (SGD/mo)"]).value == 5000
    assert ws.cell(2, COL["Salary as posted"]).value == "SGD 60,000-72,000/yr"
    g = _graded(config, good_job, salary_min=None, salary_max=None)
    ws = build_workbook([g], datetime(2026, 10, 19), {}, config["grading"]["tiers"])["Matches"]
    assert ws.cell(2, COL["Salary min (SGD/mo)"]).value is None
    assert "salary" in ws.cell(2, COL["Gaps (requirements not met)"]).value.lower()


def test_summary_uses_formulas(config, good_job):
    a = _graded(config, good_job)
    ws = build_workbook([a, a], datetime(2026, 10, 19), {}, config["grading"]["tiers"])["Summary"]
    assert ws["B4"].value == '=COUNTIF(Matches!$A$2:$A$3,"A")'
    assert ws["B7"].value == "=SUM(B4:B6)"


def test_empty_workbook_is_valid(config):
    data = workbook_bytes([], datetime(2026, 10, 19), {}, config["grading"]["tiers"])
    wb = load_workbook(BytesIO(data))
    assert wb["Matches"].max_row == 1


def test_email_attaches_xlsx(monkeypatch, config, good_job):
    sent = {}

    class FakeSMTP:
        def __init__(self, *a, **k): pass
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def ehlo(self): pass
        def starttls(self): pass
        def login(self, *a): pass
        def send_message(self, msg): sent["msg"] = msg

    monkeypatch.setenv("SMTP_USER", "me@example.com")
    monkeypatch.setenv("SMTP_PASSWORD", "x")
    monkeypatch.setattr(email_mod.smtplib, "SMTP", FakeSMTP)
    data = workbook_bytes([_graded(config, good_job)], datetime(2026, 10, 19), {}, config["grading"]["tiers"])
    email_mod.send_email("s", "<p>h</p>", "t", ["me@example.com"],
                         attachments={"digest.xlsx": data, "digest.md": "# md", "empty.xlsx": b""})
    msg = message_from_bytes(sent["msg"].as_bytes())
    parts = {p.get_filename(): p for p in msg.walk() if p.get_filename()}
    assert set(parts) == {"digest.xlsx", "digest.md"}
    assert parts["digest.xlsx"].get_content_type() == \
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert parts["digest.xlsx"].get_payload(decode=True) == data
